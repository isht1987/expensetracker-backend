from django.contrib import admin
from django.utils.html import format_html
from .models import Expense

@admin.register(Expense)
class ExpenseAdmin(admin.ModelAdmin):
    list_display = ('user', 'display_category', 'amount', 'date', 'receipt_preview', 'created_at')
    list_filter = ('category', 'date', 'user', 'created_at')
    search_fields = ('description', 'user__username', 'user__email', 'category', 'custom_category')
    date_hierarchy = 'date'
    readonly_fields = ('created_at', 'updated_at')
    ordering = ['-date', '-created_at']
    
    fieldsets = (
        ('Basic Information', {
            'fields': ('user', 'date', 'amount')
        }),
        ('Category', {
            'fields': ('category', 'custom_category'),
            'description': 'Select a category. If "Others" is selected, specify the custom category name.'
        }),
        ('Receipt', {
            'fields': ('receipt_image',)
        }),
        ('Details', {
            'fields': ('description',)
        }),
        ('Timestamps', {
            'fields': ('created_at', 'updated_at'),
            'classes': ('collapse',)
        }),
    )
    
    def display_category(self, obj):
        """Display the actual category name (including custom category for 'Others')"""
        category_name = obj.get_category_display_name()
        
        # Add a badge-like styling for custom categories
        if obj.category == 'Others' and obj.custom_category:
            return format_html(
                '<span style="background-color: #fef3c7; color: #92400e; padding: 3px 8px; border-radius: 4px; font-size: 12px;">{}</span>',
                category_name
            )
        return category_name
    
    display_category.short_description = "Category"
    display_category.admin_order_field = 'category'
    
    def receipt_preview(self, obj):
        """Display a thumbnail preview of the receipt image"""
        if obj.receipt_image:
            return format_html(
                '<a href="{}" target="_blank"><img src="{}" style="width: 50px; height: 50px; object-fit: cover; border-radius: 4px; border: 1px solid #ddd;" /></a>',
                obj.receipt_image.url,
                obj.receipt_image.url
            )
        return format_html('<span style="color: #999; font-style: italic;">No image</span>')
    
    receipt_preview.short_description = "Receipt"
    
    def save_model(self, request, obj, form, change):
        """Custom save to clear custom_category if category is not 'Others'"""
        if obj.category != 'Others':
            obj.custom_category = ''
        super().save_model(request, obj, form, change)
    
    # Add action to export selected expenses
    actions = ['export_to_excel']
    
    def export_to_excel(self, request, queryset):
        """Export selected expenses to Excel"""
        from django.http import HttpResponse
        import openpyxl
        from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
        
        # Create workbook
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Expenses"
        
        # Styling
        header_font = Font(bold=True, size=12, color="FFFFFF")
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        center_alignment = Alignment(horizontal='center', vertical='center')
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        # Headers
        headers = ['Date', 'User', 'Category', 'Amount', 'Description']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_alignment
            cell.border = border
        
        # Data rows
        for row, expense in enumerate(queryset.order_by('date'), 2):
            ws.cell(row=row, column=1, value=expense.date.strftime('%Y-%m-%d')).border = border
            ws.cell(row=row, column=2, value=expense.user.username).border = border
            ws.cell(row=row, column=3, value=expense.get_category_display_name()).border = border
            ws.cell(row=row, column=4, value=float(expense.amount)).border = border
            ws.cell(row=row, column=5, value=expense.description or "").border = border
        
        # Column widths
        ws.column_dimensions['A'].width = 12
        ws.column_dimensions['B'].width = 15
        ws.column_dimensions['C'].width = 20
        ws.column_dimensions['D'].width = 12
        ws.column_dimensions['E'].width = 30
        ws.column_dimensions['F'].width = 35
        
        # Create response
        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename="expenses_export.xlsx"'
        wb.save(response)
        
        return response
    
    export_to_excel.short_description = "Export selected expenses to Excel"