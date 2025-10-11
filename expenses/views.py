# Add this to your views.py to support quarter and half-year filtering

from rest_framework import status
from rest_framework.decorators import api_view, permission_classes
from rest_framework.permissions import IsAuthenticated
from rest_framework.pagination import PageNumberPagination
from django.db.models import Sum, Count, Q
from django.http import HttpResponse
from django.utils import timezone
from datetime import timedelta, date
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from .models import Expense
from .serializer import ExpenseSerializer, ExpenseCreateSerializer
from expenses_backend.generics import ListCreateAPIView, RetrieveUpdateDestroyAPIView
from expenses_backend.response_schemas import create_api_response
import calendar


def get_quarter_date_range(year, quarter):
    """Get start and end date for a given quarter"""
    quarter_months = {
        1: (1, 3),   # Q1: Jan-Mar
        2: (4, 6),   # Q2: Apr-Jun
        3: (7, 9),   # Q3: Jul-Sep
        4: (10, 12)  # Q4: Oct-Dec
    }
    
    start_month, end_month = quarter_months[int(quarter)]
    start_date = date(int(year), start_month, 1)
    
    # Get last day of end_month
    last_day = calendar.monthrange(int(year), end_month)[1]
    end_date = date(int(year), end_month, last_day)
    
    return start_date, end_date


def get_half_year_date_range(year, half):
    """Get start and end date for a given half year"""
    if int(half) == 1:
        # H1: Jan-Jun
        start_date = date(int(year), 1, 1)
        end_date = date(int(year), 6, 30)
    else:
        # H2: Jul-Dec
        start_date = date(int(year), 7, 1)
        end_date = date(int(year), 12, 31)
    
    return start_date, end_date


class ExpensePagination(PageNumberPagination):
    page_size = 10
    page_size_query_param = 'page_size'
    max_page_size = 100


class ExpenseListCreateView(ListCreateAPIView):
    permission_classes = [IsAuthenticated]
    pagination_class = ExpensePagination
    
    def get_queryset(self):
        queryset = Expense.objects.filter(user=self.request.user).select_related('user')
        
        # Quarter filtering
        quarter = self.request.query_params.get('quarter')
        year = self.request.query_params.get('year')
        
        if quarter and year:
            start_date, end_date = get_quarter_date_range(year, quarter)
            queryset = queryset.filter(date__gte=start_date, date__lte=end_date)
        
        # Half-year filtering
        half = self.request.query_params.get('half')
        if half and year and not quarter:
            start_date, end_date = get_half_year_date_range(year, half)
            queryset = queryset.filter(date__gte=start_date, date__lte=end_date)
        
        # Standard date filtering
        start_date_param = self.request.query_params.get('start_date')
        end_date_param = self.request.query_params.get('end_date')
        if start_date_param:
            queryset = queryset.filter(date__gte=start_date_param)
        if end_date_param:
            queryset = queryset.filter(date__lte=end_date_param)
        
        # Category filtering
        category = self.request.query_params.get('category')
        if category:
            queryset = queryset.filter(category=category)
        
        # Month/Year filtering (only if quarter/half not specified)
        month = self.request.query_params.get('month')
        if year and not quarter and not half:
            queryset = queryset.filter(date__year=year)
        if month and year and not quarter and not half:
            queryset = queryset.filter(date__month=month)
        
        # Search functionality
        search = self.request.query_params.get('search')
        if search:
            queryset = queryset.filter(
                Q(description__icontains=search) |
                Q(category__icontains=search) |
                Q(custom_category__icontains=search)
            )
        
        return queryset.order_by('-date', '-created_at')
    
    def get_serializer_class(self):
        if self.request.method == 'POST':
            return ExpenseCreateSerializer
        return ExpenseSerializer
    
    def perform_create(self, serializer):
        serializer.save(user=self.request.user)


class ExpenseDetailView(RetrieveUpdateDestroyAPIView):
    serializer_class = ExpenseSerializer
    permission_classes = [IsAuthenticated]
    
    def get_queryset(self):
        return Expense.objects.filter(user=self.request.user)
    
    def get_serializer_class(self):
        if self.request.method in ['PUT', 'PATCH']:
            return ExpenseCreateSerializer
        return ExpenseSerializer


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def export_excel(request):
    """Export expenses to Excel with filtering - supports quarter and half-year"""
    year = request.GET.get('year')
    month = request.GET.get('month')
    quarter = request.GET.get('quarter')  # NEW: 1, 2, 3, or 4
    half = request.GET.get('half')  # NEW: 1 or 2
    start_date = request.GET.get('start_date')
    end_date = request.GET.get('end_date')
    category = request.GET.get('category')
    
    queryset = Expense.objects.filter(user=request.user)
    
    # Filename generation
    filename_parts = ["expenses"]
    filter_description = []
    
    # Apply filters with priority: custom range > quarter > half > month > year
    if start_date and end_date:
        queryset = queryset.filter(date__gte=start_date, date__lte=end_date)
        filename_parts.append(f"{start_date}_to_{end_date}")
        filter_description.append(f"Period: {start_date} to {end_date}")
    
    elif quarter and year:
        start_date_q, end_date_q = get_quarter_date_range(year, quarter)
        queryset = queryset.filter(date__gte=start_date_q, date__lte=end_date_q)
        filename_parts.extend([f"Q{quarter}", year])
        filter_description.append(f"Quarter: Q{quarter} {year}")
    
    elif half and year:
        start_date_h, end_date_h = get_half_year_date_range(year, half)
        queryset = queryset.filter(date__gte=start_date_h, date__lte=end_date_h)
        filename_parts.extend([f"H{half}", year])
        filter_description.append(f"Half: H{half} {year}")
    
    elif year and month:
        queryset = queryset.filter(date__year=year, date__month=month)
        month_name = calendar.month_name[int(month)]
        filename_parts.extend([month_name, year])
        filter_description.append(f"Month: {month_name} {year}")
    
    elif year:
        queryset = queryset.filter(date__year=year)
        filename_parts.append(year)
        filter_description.append(f"Year: {year}")
    
    else:
        # Default to current year
        current_year = timezone.now().year
        queryset = queryset.filter(date__year=current_year)
        filename_parts.append(str(current_year))
        filter_description.append(f"Year: {current_year}")
    
    if category:
        queryset = queryset.filter(category=category)
        filename_parts.append(category.replace(' ', '_'))
        filter_description.append(f"Category: {category}")
    
    filename_parts.append(request.user.username)
    filename = "_".join(filename_parts)
    
    # Order by date
    queryset = queryset.order_by('date', 'created_at')
    
    # Create Excel workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Expenses"
    
    # Styling
    header_font = Font(bold=True, size=12, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    center_alignment = Alignment(horizontal='center', vertical='center')
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # Title row
    ws.merge_cells('A1:E1')
    title_cell = ws['A1']
    title_cell.value = f"EXPENSE REPORT - {request.user.get_full_name() or request.user.username}"
    title_cell.font = Font(bold=True, size=14)
    title_cell.alignment = center_alignment
    
    # Filter info row
    ws.merge_cells('A2:E2')
    filter_cell = ws['A2']
    filter_cell.value = " | ".join(filter_description)
    filter_cell.font = Font(size=10, italic=True)
    filter_cell.alignment = center_alignment
    
    # Headers
    headers = ['Date', 'Category', 'Amount', 'Description', 'Created At']
    
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_alignment
        cell.border = border
    
    # Data rows
    # Data rows
    for row, expense in enumerate(queryset, 5):
        ws.cell(row=row, column=1, value=expense.date.strftime('%Y-%m-%d')).border = border
        ws.cell(row=row, column=2, value=expense.get_category_display_name()).border = border
        ws.cell(row=row, column=3, value=float(expense.amount)).border = border
        ws.cell(row=row, column=4, value=expense.description or "").border = border
        ws.cell(row=row, column=5, value=expense.created_at.strftime('%Y-%m-%d %H:%M')).border = border  # ✅ NOW INSIDE THE LOOP
        
    # Summary section
    summary_row = queryset.count() + 6
    ws.cell(row=summary_row, column=1, value="TOTAL:").font = Font(bold=True, size=12)
    total_amount = queryset.aggregate(Sum('amount'))['amount__sum'] or 0
    total_cell = ws.cell(row=summary_row, column=3, value=float(total_amount))
    total_cell.font = Font(bold=True, size=12)
    total_cell.fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
    
    ws.cell(row=summary_row + 1, column=1, value="Total Expenses:").font = Font(bold=True)
    ws.cell(row=summary_row + 1, column=3, value=queryset.count()).font = Font(bold=True)
    
    # Column widths
    column_widths = [12, 20, 12, 30, 18]
    for col, width in enumerate(column_widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = width
    
    # Create HTTP response
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}.xlsx"'
    
    wb.save(response)
    return response


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def dashboard_stats(request):
    """Get dashboard statistics"""
    user_expenses = Expense.objects.filter(user=request.user)
    
    now = timezone.now()
    current_month = user_expenses.filter(
        date__year=now.year,
        date__month=now.month
    )
    
    current_year = user_expenses.filter(date__year=now.year)
    
    # Previous month for comparison
    first_day_current_month = now.replace(day=1)
    last_day_prev_month = first_day_current_month - timedelta(days=1)
    previous_month = user_expenses.filter(
        date__year=last_day_prev_month.year,
        date__month=last_day_prev_month.month
    )
    
    # Category breakdown for current month
    category_data = []
    for expense in current_month:
        category_name = expense.get_category_display_name()
        existing = next((item for item in category_data if item['category'] == category_name), None)
        if existing:
            existing['total'] += float(expense.amount)
            existing['count'] += 1
        else:
            category_data.append({
                'category': category_name,
                'total': float(expense.amount),
                'count': 1
            })
    
    category_data.sort(key=lambda x: x['total'], reverse=True)
    
    # Monthly trend for current year
    monthly_trend = []
    for month_num in range(1, 13):
        month_expenses = user_expenses.filter(
            date__year=now.year,
            date__month=month_num
        ).aggregate(total=Sum('amount'))
        monthly_trend.append({
            'month': calendar.month_name[month_num],
            'month_number': month_num,
            'total': float(month_expenses['total'] or 0)
        })
    
    # Top categories for current year
    year_category_data = []
    for expense in current_year:
        category_name = expense.get_category_display_name()
        existing = next((item for item in year_category_data if item['category'] == category_name), None)
        if existing:
            existing['total'] += float(expense.amount)
        else:
            year_category_data.append({
                'category': category_name,
                'total': float(expense.amount)
            })
    
    year_category_data.sort(key=lambda x: x['total'], reverse=True)
    
    stats = {
        'current_month': {
            'total': float(current_month.aggregate(Sum('amount'))['amount__sum'] or 0),
            'count': current_month.count(),
            'month_name': calendar.month_name[now.month]
        },
        'current_year': {
            'total': float(current_year.aggregate(Sum('amount'))['amount__sum'] or 0),
            'count': current_year.count(),
            'year': now.year
        },
        'previous_month': {
            'total': float(previous_month.aggregate(Sum('amount'))['amount__sum'] or 0),
            'count': previous_month.count(),
            'month_name': calendar.month_name[last_day_prev_month.month]
        },
        'category_breakdown': category_data,
        'monthly_trend': monthly_trend,
        'recent_expenses': ExpenseSerializer(
            user_expenses.order_by('-date', '-created_at')[:5], 
            many=True, 
            context={'request': request}
        ).data,
        'top_categories': year_category_data[:5]
    }
    
    return create_api_response(
        request,
        status_code=status.HTTP_200_OK,
        message="Dashboard stats retrieved successfully",
        data=stats
    )


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def expense_history(request):
    """Get expense history with filtering options"""
    queryset = Expense.objects.filter(user=request.user)
    
    # Apply filters
    year = request.GET.get('year')
    month = request.GET.get('month')
    quarter = request.GET.get('quarter')
    half = request.GET.get('half')
    category = request.GET.get('category')
    start_date = request.GET.get('start_date')
    end_date = request.GET.get('end_date')
    search = request.GET.get('search')
    
    # Priority filtering
    if start_date and end_date:
        queryset = queryset.filter(date__gte=start_date, date__lte=end_date)
    elif quarter and year:
        start_q, end_q = get_quarter_date_range(year, quarter)
        queryset = queryset.filter(date__gte=start_q, date__lte=end_q)
    elif half and year:
        start_h, end_h = get_half_year_date_range(year, half)
        queryset = queryset.filter(date__gte=start_h, date__lte=end_h)
    elif year and month:
        queryset = queryset.filter(date__year=year, date__month=month)
    elif year:
        queryset = queryset.filter(date__year=year)
    
    if category:
        queryset = queryset.filter(category=category)
    if search:
        queryset = queryset.filter(
            Q(description__icontains=search) |
            Q(category__icontains=search) |
            Q(custom_category__icontains=search)
        )
    
    # Pagination
    paginator = ExpensePagination()
    paginated_queryset = paginator.paginate_queryset(queryset.order_by('-date', '-created_at'), request)
    
    serializer = ExpenseSerializer(paginated_queryset, many=True, context={'request': request})
    
    # Calculate totals
    total_amount = queryset.aggregate(Sum('amount'))['amount__sum'] or 0
    total_count = queryset.count()
    
    return paginator.get_paginated_response({
        'expenses': serializer.data,
        'summary': {
            'total_amount': float(total_amount),
            'total_count': total_count
        }
    })


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def expense_summary(request):
    """Get expense summary statistics"""
    year = request.GET.get('year', timezone.now().year)
    month = request.GET.get('month')
    quarter = request.GET.get('quarter')
    half = request.GET.get('half')
    
    queryset = Expense.objects.filter(user=request.user, date__year=year)
    
    period_description = str(year)
    
    if quarter:
        start_q, end_q = get_quarter_date_range(year, quarter)
        queryset = queryset.filter(date__gte=start_q, date__lte=end_q)
        period_description = f"Q{quarter} {year}"
    elif half:
        start_h, end_h = get_half_year_date_range(year, half)
        queryset = queryset.filter(date__gte=start_h, date__lte=end_h)
        period_description = f"H{half} {year}"
    elif month:
        queryset = queryset.filter(date__month=month)
        period_description = f"{calendar.month_name[int(month)]} {year}"
    
    # Category-wise summary
    category_summary_data = []
    for expense in queryset:
        category_name = expense.get_category_display_name()
        existing = next((item for item in category_summary_data if item['category'] == category_name), None)
        if existing:
            existing['total'] += float(expense.amount)
            existing['count'] += 1
        else:
            category_summary_data.append({
                'category': category_name,
                'total': float(expense.amount),
                'count': 1
            })
    
    # Calculate average
    for item in category_summary_data:
        item['avg'] = item['total'] / item['count']
    
    category_summary_data.sort(key=lambda x: x['total'], reverse=True)
    
    # Month-wise summary (if year is selected without month/quarter/half)
    monthly_summary = []
    if not month and not quarter and not half:
        for m in range(1, 13):
            month_data = queryset.filter(date__month=m).aggregate(
                total=Sum('amount'),
                count=Count('id')
            )
            monthly_summary.append({
                'month': calendar.month_name[m],
                'month_number': m,
                'total': float(month_data['total'] or 0),
                'count': month_data['count'] or 0
            })
    
    total = queryset.aggregate(Sum('amount'))['amount__sum'] or 0
    
    summary = {
        'total_amount': float(total),
        'total_count': queryset.count(),
        'period': period_description,
        'category_summary': category_summary_data,
        'monthly_summary': monthly_summary
    }
    
    return create_api_response(
        request,
        status_code=status.HTTP_200_OK,
        message="Expense summary retrieved successfully",
        data=summary
    )


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def category_choices(request):
    """Get available category choices"""
    categories = [choice[0] for choice in Expense.CATEGORY_CHOICES]
    return create_api_response(
        request,
        status_code=status.HTTP_200_OK,
        message="Categories retrieved successfully",
        data={'categories': categories}
    )